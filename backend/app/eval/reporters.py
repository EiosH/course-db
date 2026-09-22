"""Pluggable eval reporters: Excel / plain-text (and anything else you plug in)."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Protocol

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from answering import format_hits, plan_resolve
from config import OUTPUT_DIR
from pipeline import QueryResult

ANSWER_HEADERS = [
    "序号",
    "原问题",
    "查询计划",
    "预探查",
    "预探查召回",
    "rewritten_query",
    "hard_constraints",
    "course_id",
    "quarter",
    "lecturer",
    "screen_shot语义召回",
    "screen_shot关键词召回",
    "transcript语义召回",
    "transcript关键词召回",
    "rerank后检索",
    "答案",
]
# 1-based Excel column of the first answer cell ("答案")
ANSWER_COL = ANSWER_HEADERS.index("答案") + 1
QUERY_COL = ANSWER_HEADERS.index("原问题") + 1
ANSWER_COL_WIDTHS = {
    "A": 6,
    "B": 36,
    "C": 36,
    "D": 36,
    "E": 42,
    "F": 40,
    "G": 28,
    "H": 12,
    "I": 14,
    "J": 20,
    "K": 42,
    "L": 42,
    "M": 42,
    "N": 42,
    "O": 50,
    "P": 50,
}
DEFAULT_ANSWER_WIDTH = 50
CELL_WRAP = Alignment(vertical="top", wrap_text=True)


def format_query_plan(preprobe: dict | None) -> str:
    """Query plan for Excel / logs."""
    if not preprobe:
        return ""
    plan = preprobe.get("plan") or {}
    return json.dumps(
        {
            "time_mode": plan.get("time_mode"),
            "hard_constraints": plan.get("hard_constraints") or [],
            "probe_hard_constraints": plan.get("probe_hard_constraints") or [],
            "course_general_knowledge": bool(plan.get("course_general_knowledge")),
            "resolve": plan_resolve(plan),
            "reason": plan.get("reason") or "",
        },
        ensure_ascii=False,
        indent=2,
    )


def format_preprobe(preprobe: dict | None) -> str:
    """Resolve-search execution summary."""
    if not preprobe:
        return ""
    plan = preprobe.get("plan") or {}
    target = plan_resolve(plan)
    if not target:
        return "skipped"
    resolved = preprobe.get("resolved") or {}
    conf = resolved.get("confidence")
    conf_s = f"{conf:.2f}" if isinstance(conf, (int, float)) else "?"
    return "\n".join(
        [
            f"resolve={target}",
            f"query={preprobe.get('query') or ''}",
            f"hits={len(preprobe.get('hits') or [])}",
            f"found={bool(resolved.get('found'))}",
            f"name={resolved.get('name') or ''}",
            f"confidence={conf_s}",
        ]
    )


class Reporter(Protocol):
    """Eval sink: start → record each QueryResult → finish."""

    def start(self) -> None: ...

    def record(self, index: int, result: QueryResult) -> None: ...

    def finish(self) -> None: ...


class ExcelReporter:
    """
    Persistent .xlsx dump of retrieval + answer (eval-only).

    Same question → append answer in the next empty answer column.
    New question → append a full new row.
    """

    def __init__(self, path: Path | str | None = None):
        self.path = Path(path) if path else OUTPUT_DIR / "answer.xlsx"
        self.wb = None
        self.ws = None

    def start(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if self.path.exists():
            self.wb = load_workbook(self.path)
            self.ws = self.wb.active
            print(f"[excel] appending → {self.path}")
            return
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "answers"
        self.ws.append(ANSWER_HEADERS)
        for cell in self.ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = CELL_WRAP
        for col, width in ANSWER_COL_WIDTHS.items():
            self.ws.column_dimensions[col].width = width
        self.wb.save(self.path)
        print(f"[excel] writing → {self.path}")

    def _find_question_row(self, query: str) -> int | None:
        for row_idx in range(2, (self.ws.max_row or 1) + 1):
            val = self.ws.cell(row=row_idx, column=QUERY_COL).value
            if val is not None and str(val).strip() == query.strip():
                return row_idx
        return None

    def _next_answer_col(self, row_idx: int) -> int:
        col = ANSWER_COL
        while True:
            val = self.ws.cell(row=row_idx, column=col).value
            if val is None or (isinstance(val, str) and not val.strip()):
                return col
            col += 1

    def _ensure_answer_header(self, col: int) -> None:
        cell = self.ws.cell(row=1, column=col)
        if cell.value:
            return
        n = col - ANSWER_COL + 1
        cell.value = "答案" if n == 1 else f"答案{n}"
        cell.font = Font(bold=True)
        cell.alignment = CELL_WRAP
        letter = get_column_letter(col)
        if not self.ws.column_dimensions[letter].width:
            self.ws.column_dimensions[letter].width = DEFAULT_ANSWER_WIDTH

    def _append_answer_cell(self, row_idx: int, answer_text: str) -> int:
        col = self._next_answer_col(row_idx)
        self._ensure_answer_header(col)
        cell = self.ws.cell(row=row_idx, column=col, value=answer_text)
        cell.alignment = CELL_WRAP
        return col

    def record(self, index: int, result: QueryResult) -> None:
        existing = self._find_question_row(result.query)
        if existing is not None:
            col = self._append_answer_cell(existing, result.answer_text)
            self.wb.save(self.path)
            letter = get_column_letter(col)
            print(f"[excel] same question → row {existing} col {letter} → {self.path}")
            return

        channel = "time" if result.ts_value else "semantic"
        final_channel = "time" if result.ts_value else "rerank"
        preprobe_hits = (result.preprobe or {}).get("hits") or []
        row = [
            index,
            result.query,
            format_query_plan(result.preprobe),
            format_preprobe(result.preprobe),
            format_hits(preprobe_hits, "preprobe"),
            result.search_query,
            json.dumps(
                result.rewritten.get("hard_constraints", []), ensure_ascii=False
            ),
            (result.course_ctx or {}).get("course_id", ""),
            (result.course_ctx or {}).get("quarter", ""),
            (result.course_ctx or {}).get("lecturer", ""),
            format_hits(result.ss_dense, channel),
            format_hits(result.ss_bm25, "keyword"),
            format_hits(result.tr_dense, channel),
            format_hits(result.tr_bm25, "keyword"),
            format_hits(result.final_hits, final_channel),
            result.answer_text,
        ]
        self.ws.append(row)
        for cell in self.ws[self.ws.max_row]:
            cell.alignment = CELL_WRAP
        self.wb.save(self.path)
        print(f"[excel] new row {self.ws.max_row} → {self.path}")

    def finish(self) -> None:
        print(f"[excel] done: {self.path}")


class TxtReporter:
    """Human-readable .txt log (eval-only), similar to legacy answer.txt."""

    def __init__(self, path: Path | str | None = None):
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        self.path = Path(path) if path else OUTPUT_DIR / f"answer_{stamp}.txt"
        self._fh = None

    def start(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._fh = open(self.path, "w", encoding="utf-8")
        print(f"[txt] writing → {self.path}")

    def record(self, index: int, result: QueryResult) -> None:
        preprobe_hits = (result.preprobe or {}).get("hits") or []
        hits_block = format_hits(preprobe_hits, "preprobe") or "(none)"
        block = (
            f"{index}. 原问题: {result.query}\n"
            f"query_plan:\n"
            f"{format_query_plan(result.preprobe)}\n"
            f"\n"
            f"preprobe:\n"
            f"{format_preprobe(result.preprobe)}\n"
            f"\n"
            f"preprobe hits:\n"
            f"{hits_block}\n"
            f"\n"
            f"rewrite:\n"
            f"{json.dumps(result.rewritten, ensure_ascii=False, indent=2)}\n"
            f"\n"
            f"prompt:\n"
            f"{result.prompt}\n"
            f"\n"
            f"答案:\n"
            f"{result.answer_text}\n"
            f"\n"
            f"{'=' * 60}\n\n"
        )
        self._fh.write(block)
        self._fh.flush()
        print(f"[txt] wrote item {index} → {self.path}")

    def finish(self) -> None:
        if self._fh:
            self._fh.close()
            self._fh = None
        print(f"[txt] done: {self.path}")


class ConsoleReporter:
    """Print rewrite / prompt / answer to stdout (always useful while debugging)."""

    def start(self) -> None:
        return

    def record(self, index: int, result: QueryResult) -> None:
        print(f"\nquery:    {result.query}")
        ctx = result.course_ctx or {}
        if ctx:
            print(
                f"course:   {ctx.get('course_id')} / {ctx.get('quarter')} / "
                f"{ctx.get('lecturer')} (lecture_id={ctx.get('lecture_id')})"
            )
        print("query_plan:")
        print(format_query_plan(result.preprobe) or "(none)")
        print("preprobe:")
        print(format_preprobe(result.preprobe) or "(none)")
        preprobe_hits = (result.preprobe or {}).get("hits") or []
        if preprobe_hits:
            print(f"preprobe hits ({len(preprobe_hits)}):")
            print(format_hits(preprobe_hits, "preprobe"))
        print("rewrite:")
        print(json.dumps(result.rewritten, ensure_ascii=False, indent=2))
        if result.search_query != result.rewritten.get("rewritten_query"):
            print("rewrite (after anchor expand):")
            q = result.search_query
            print(q[:500] + ("..." if len(q) > 500 else ""))
        n_ss = sum(
            1 for h in result.final_hits if h.payload.get("type") == "screen_shot"
        )
        n_tr = sum(
            1 for h in result.final_hits if h.payload.get("type") == "transcript"
        )
        if result.ts_value:
            print(
                f"time kept {len(result.final_hits)} "
                f"(ss={n_ss} tr={n_tr}, no rerank)"
            )
            if not result.final_hits:
                print(
                    "warning: timestamp filter matched 0 chunks — "
                    "re-run with --ingest to rebuild start_sec/end_sec indexes"
                )
        else:
            n_cand = (
                len(result.ss_dense)
                + len(result.ss_bm25)
                + len(result.tr_dense)
                + len(result.tr_bm25)
            )
            print(
                f"recall screen_shot dense={len(result.ss_dense)} "
                f"bm25={len(result.ss_bm25)} | "
                f"transcript dense={len(result.tr_dense)} "
                f"bm25={len(result.tr_bm25)}"
            )
            print(f"rerank kept {len(result.final_hits)}/{n_cand}")
        print(result.prompt)
        print(f"\nanswer:\n{result.answer_text}")

    def finish(self) -> None:
        return


def default_reporters(*, excel: bool = True, txt: bool = True) -> list[Reporter]:
    reps: list[Reporter] = [ConsoleReporter()]
    if excel:
        reps.append(ExcelReporter())
    if txt:
        reps.append(TxtReporter())
    return reps
